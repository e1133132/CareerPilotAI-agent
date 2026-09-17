from __future__ import annotations

from pathlib import Path
import unicodedata


def _clean_extracted_text(text: str) -> str:
    if not text:
        return ""
    cleaned_chars: list[str] = []
    for ch in text:
        if unicodedata.category(ch) == "Cc" and ch not in ("\t", "\n", "\r"):
            continue
        cleaned_chars.append(ch)
    out = unicodedata.normalize("NFKC", "".join(cleaned_chars)).strip()
    if not out:
        return out
    try:
        from ftfy import fix_text  # type: ignore

        out = fix_text(out).strip()
    except Exception:
        pass
    return out


def _ocr_pdf_text(pdf_bytes: bytes) -> str:
    try:
        import numpy as np  # type: ignore
        import pypdfium2 as pdfium  # type: ignore
        from rapidocr_onnxruntime import RapidOCR  # type: ignore
    except ModuleNotFoundError as e:
        raise ModuleNotFoundError(
            "OCR fallback requires `pypdfium2` and `rapidocr-onnxruntime`. "
            "Please install them to parse scanned PDFs."
        ) from e

    pdf = pdfium.PdfDocument(pdf_bytes)
    ocr_engine = RapidOCR()
    lines: list[str] = []
    for i in range(len(pdf)):
        page = pdf[i]
        try:
            bitmap = page.render(scale=2.0)
            pil_img = bitmap.to_pil()
            arr = np.array(pil_img)
            result, _ = ocr_engine(arr)
            if result:
                page_text = "\n".join(str(item[1]) for item in result if len(item) >= 2 and str(item[1]).strip())
                if page_text.strip():
                    lines.append(page_text.strip())
        except Exception:
            continue
    return _clean_extracted_text("\n\n".join(lines).strip())


def load_resume_text(path: str) -> str:
    """
    Minimal resume loader for workshop use.

    Supports:
    - Plain text files (UTF-8 / system default fallback)
    - PDF files (extracts text via `pypdf`)
    - DOCX files (extracts paragraph text via `python-docx`)
    """
    p = Path(path).expanduser()
    if not p.exists():
        raise FileNotFoundError(f"Resume file not found: {p}")
    if p.is_dir():
        raise IsADirectoryError(f"Resume path is a directory: {p}")

    suffix = p.suffix.lower()
    if suffix == ".pdf":
        try:
            from io import BytesIO

            from pypdf import PdfReader  # type: ignore
        except ModuleNotFoundError as e:
            raise ModuleNotFoundError(
                "PDF support requires `pypdf`. Please install it: `pip install pypdf`."
            ) from e

        pdf_bytes = p.read_bytes()
        reader = PdfReader(BytesIO(pdf_bytes))

        pages_text: list[str] = []
        for page in reader.pages:
            try:
                t = page.extract_text() or ""
            except Exception:
                t = ""
            t = t.strip()
            if t:
                pages_text.append(t)

        resume_text = _clean_extracted_text("\n".join(pages_text).strip())
        if not resume_text:
            resume_text = _ocr_pdf_text(pdf_bytes)
        if not resume_text:
            raise ValueError("No extractable text found in the PDF, including OCR fallback.")
        return resume_text

    if suffix == ".docx":
        try:
            from io import BytesIO

            from docx import Document  # type: ignore
        except ModuleNotFoundError as e:
            raise ModuleNotFoundError(
                "DOCX support requires `python-docx`. Please install it: `pip install python-docx`."
            ) from e

        doc = Document(BytesIO(p.read_bytes()))
        paragraphs = [para.text.strip() for para in doc.paragraphs if para.text and para.text.strip()]
        resume_text = _clean_extracted_text("\n\n".join(paragraphs).strip())
        if not resume_text:
            raise ValueError("No extractable text found in the DOCX file.")
        return resume_text

    try:
        return _clean_extracted_text(p.read_text(encoding="utf-8"))
    except UnicodeDecodeError:
        return _clean_extracted_text(p.read_text())
