from __future__ import annotations

from io import BytesIO
import json
import queue
import threading
import unicodedata
import uuid
from collections.abc import Iterator
from datetime import datetime, timezone
from typing import Any

from dotenv import load_dotenv
from fastapi import BackgroundTasks, Depends, FastAPI, File, Form, Header, HTTPException, Query, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response, StreamingResponse
from pydantic import BaseModel, Field
from pypdf import PdfReader  # type: ignore

from config import settings
from agents.llm_utils import get_embed_fn
from db import init_database
from pipeline.runner import (
    finish_run_background,
    recompute_skill_gaps,
    response_payload_from_state,
    run_sync,
    start_partial_run,
)
from pipeline.session import RunSession, run_store
from security.input_guard import normalize_target_roles, validate_api_user_inputs
from skills.explainability import build_explainability_block
from skills.vector_store import warmup_qdrant_indexes
from skills.user_memory import derive_user_id, load_user_memory, save_user_memory
from skills.application_pack import (
    build_application_pack,
    pack_to_markdown,
    pack_to_zip_bytes,
    zip_download_filename,
)
from skills.auth import authenticate_user, create_access_token, register_user, verify_access_token
from workflow_graph import latest_step_message
from observability import (
    metrics_store,
    RequestMetricsMiddleware,
    set_run_id,
    setup_logging,
    setup_tracing,
)

load_dotenv(override=True)
setup_logging()

app = FastAPI(title=settings.APP_NAME)

app.add_middleware(RequestMetricsMiddleware)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


class SelectJobRequest(BaseModel):
    job_id: str = Field(min_length=1)


class RejectJobRequest(BaseModel):
    job_id: str = Field(min_length=1)


class UserMemoryPatch(BaseModel):
    preferred_target_roles: list[str] | None = None
    rejected_job_ids: list[str] | None = None
    saved_job_ids: list[str] | None = None


class RegisterRequest(BaseModel):
    username: str = Field(min_length=3, max_length=64)
    password: str = Field(min_length=6, max_length=128)
    email: str | None = None


class LoginRequest(BaseModel):
    username: str = Field(min_length=3, max_length=64)
    password: str = Field(min_length=6, max_length=128)


def _optional_bearer_token(authorization: str | None = Header(default=None)) -> dict[str, Any] | None:
    if not settings.AUTH_ENABLED:
        return None
    raw = (authorization or "").strip()
    if not raw.lower().startswith("bearer "):
        if settings.AUTH_REQUIRED:
            raise HTTPException(status_code=401, detail="Authorization Bearer token required")
        return None
    token = raw.split(" ", 1)[1].strip()
    try:
        return verify_access_token(token)
    except Exception as exc:
        if settings.AUTH_REQUIRED:
            raise HTTPException(status_code=401, detail="Invalid or expired token") from exc
        return None


def _resolve_user_id(
    *,
    resume_text: str,
    client_id: str | None,
    account: dict[str, Any] | None,
) -> tuple[str, str | None]:
    account_uid = str(account.get("user_id") or "") if account else None
    username = str(account.get("username") or "") if account else None
    user_id = derive_user_id(
        client_id=client_id or None,
        resume_text=resume_text,
        account_user_id=account_uid,
    )
    return user_id, username


def _background_warmup_qdrant() -> None:
    try:
        warmup_qdrant_indexes(get_embed_fn())
    except Exception:
        return


@app.on_event("startup")
def _startup_warmup() -> None:
    setup_tracing()
    if settings.DATABASE_AUTO_CREATE_TABLES:
        init_database()
    threading.Thread(target=_background_warmup_qdrant, daemon=True).start()


def _extract_pdf_text(raw_bytes: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(raw_bytes))
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to read PDF: {str(e)}") from e
    pages_text: list[str] = []
    for page in reader.pages:
        try:
            t = (page.extract_text() or "").strip()
        except Exception:
            t = ""
        if t:
            pages_text.append(t)
    return _clean_extracted_text("\n\n".join(pages_text).strip())


def _extract_docx_text(raw_bytes: bytes) -> str:
    try:
        from docx import Document  # type: ignore
    except ModuleNotFoundError as e:
        raise HTTPException(
            status_code=500,
            detail="DOCX support requires `python-docx`.",
        ) from e

    try:
        doc = Document(BytesIO(raw_bytes))
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to read DOCX: {str(e)}") from e

    lines = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    return _clean_extracted_text("\n\n".join(lines).strip())


def _clean_extracted_text(text: str) -> str:
    """
    Best-effort text cleanup:
    - remove unsafe control chars
    - normalize unicode form
    - optional mojibake repair via ftfy
    """
    if not text:
        return ""
    cleaned_chars: list[str] = []
    for ch in text:
        if unicodedata.category(ch) == "Cc" and ch not in ("\t", "\n", "\r"):
            continue
        cleaned_chars.append(ch)
    out = unicodedata.normalize("NFKC", "".join(cleaned_chars)).strip()
    if not out:
        return out
    try:
        from ftfy import fix_text  # type: ignore

        out = fix_text(out).strip()
    except Exception:
        # ftfy is optional; skip if unavailable
        pass
    return out


def _ocr_pdf_text(raw_bytes: bytes) -> str:
    """
    OCR fallback for scanned PDFs (no extractable text layer).
    Uses pypdfium2 for page rendering + RapidOCR for recognition.
    """
    try:
        import numpy as np  # type: ignore
        import pypdfium2 as pdfium  # type: ignore
        from rapidocr_onnxruntime import RapidOCR  # type: ignore
    except ModuleNotFoundError as e:
        raise HTTPException(
            status_code=500,
            detail=(
                "OCR fallback requires `pypdfium2` and `rapidocr-onnxruntime`. "
                "Install them to enable scanned PDF support."
            ),
        ) from e

    try:
        pdf = pdfium.PdfDocument(raw_bytes)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to read PDF for OCR: {str(e)}") from e

    ocr_engine = RapidOCR()
    lines: list[str] = []
    for i in range(len(pdf)):
        page = pdf[i]
        try:
            # Scale≈2.0 gives readable OCR quality for most resumes.
            bitmap = page.render(scale=2.0)
            pil_img = bitmap.to_pil()
            arr = np.array(pil_img)
            result, _ = ocr_engine(arr)
            if result:
                page_text = "\n".join(str(item[1]) for item in result if len(item) >= 2 and str(item[1]).strip())
                if page_text.strip():
                    lines.append(page_text.strip())
        except Exception:
            continue

    return _clean_extracted_text("\n\n".join(lines).strip())


def _extract_resume_text(raw_bytes: bytes, filename: str | None) -> str:
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        text = _extract_pdf_text(raw_bytes)
        if text:
            return text
        # Scanned PDF often has no text layer -> OCR fallback.
        return _ocr_pdf_text(raw_bytes)
    if name.endswith(".docx"):
        return _extract_docx_text(raw_bytes)
    raise HTTPException(status_code=400, detail="Only .pdf and .docx resume files are supported")


async def _read_resume_upload(resume_file: UploadFile) -> tuple[str, bytes]:
    if resume_file.filename and not resume_file.filename.lower().endswith((".pdf", ".docx")):
        raise HTTPException(status_code=400, detail="Only .pdf and .docx resume files are supported")
    raw_bytes = await resume_file.read()
    if not raw_bytes:
        raise HTTPException(status_code=400, detail="Uploaded resume_file is empty")
    resume_text = _extract_resume_text(raw_bytes, resume_file.filename)
    if not resume_text:
        raise HTTPException(
            status_code=400,
            detail="Could not extract text from the uploaded file. For scanned PDFs/images, OCR is required.",
        )
    return resume_text, raw_bytes


@app.get("/health")
def health() -> dict:
    from observability.tracing import configure_langsmith_env

    tracing = configure_langsmith_env()
    return {
        "ok": True,
        "service": settings.APP_NAME,
        "tracing": {
            "langsmith_enabled": tracing["langsmith_enabled"],
            "langsmith_project": tracing["langsmith_project"],
            "callbacks_background": tracing["langsmith_callbacks_background"],
        },
    }


@app.get("/metrics")
def metrics() -> dict:
    if not settings.METRICS_ENDPOINT_ENABLED:
        raise HTTPException(status_code=404, detail="metrics endpoint disabled")
    return {"ok": True, "metrics": metrics_store.snapshot()}


@app.post("/api/auth/register")
def auth_register(body: RegisterRequest) -> dict:
    if not settings.AUTH_ENABLED:
        raise HTTPException(status_code=503, detail="Auth is disabled")
    try:
        user = register_user(username=body.username, password=body.password, email=body.email)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    token = create_access_token(user_id=user["user_id"], username=user["username"])
    return {"ok": True, "token": token, "user": user}


@app.post("/api/auth/login")
def auth_login(body: LoginRequest) -> dict:
    if not settings.AUTH_ENABLED:
        raise HTTPException(status_code=503, detail="Auth is disabled")
    try:
        user = authenticate_user(username=body.username, password=body.password)
    except ValueError as exc:
        raise HTTPException(status_code=401, detail=str(exc)) from exc
    token = create_access_token(user_id=user["user_id"], username=user["username"])
    return {"ok": True, "token": token, "user": user}


@app.get("/api/auth/me")
def auth_me(account: dict[str, Any] | None = Depends(_optional_bearer_token)) -> dict:
    if not account:
        raise HTTPException(status_code=401, detail="Not authenticated")
    mem = load_user_memory(str(account["user_id"]))
    return {"ok": True, "user": account, "user_memory": mem}


@app.post("/api/careerpilot/run")
async def run_careerpilot(
    resume_file: UploadFile = File(...),
    target_roles: str = Form(default=""),
) -> dict:
    resume_text, _ = await _read_resume_upload(resume_file)
    roles = normalize_target_roles(target_roles)
    resume_text = validate_api_user_inputs(resume_text, roles)

    result = run_sync(
        initial_state={
            "resume_text": resume_text,
            "resume_path": None,
            "target_roles": roles or None,
        }
    )
    state = result.get("state") or {}

    return {
        "ok": True,
        "candidate_profile": state.get("candidate_profile"),
        "resume_evidence": state.get("resume_evidence"),
        "recommended_jobs": state.get("job_matches"),
        "skill_gaps": state.get("skill_gaps"),
        "study_plan": state.get("study_plan"),
        "resume_suggestions": state.get("resume_suggestions"),
        "apply_strategy": state.get("apply_strategy"),
        "application_pack": state.get("application_pack"),
        "report_text": result.get("report_text") or "",
        "explainability": build_explainability_block(state),
        "full_state": state,
    }


@app.post("/api/careerpilot/run_partial")
async def run_careerpilot_partial(
    background_tasks: BackgroundTasks,
    resume_file: UploadFile = File(...),
    target_roles: str = Form(default=""),
    client_id: str = Form(default=""),
    account: dict[str, Any] | None = Depends(_optional_bearer_token),
) -> dict:
    """
    Fast path: returns candidate_profile/job_matches/skill_gaps after phase-1 agents,
    then continues phase-2 via LangGraph interrupt steps in the background.
    """
    resume_text, _ = await _read_resume_upload(resume_file)
    roles = normalize_target_roles(target_roles)
    resume_text = validate_api_user_inputs(resume_text, roles)

    user_id, username = _resolve_user_id(resume_text=resume_text, client_id=client_id or None, account=account)
    user_memory = load_user_memory(user_id)

    run_id = uuid.uuid4().hex
    initial_state: dict[str, Any] = {
        "resume_text": resume_text,
        "resume_path": None,
        "target_roles": roles or None,
        "user_id": user_id,
        "user_memory": user_memory,
        "run_id": run_id,
        "stage": "intake",
        "messages": [],
        "pipeline_trace": [],
        "fallback_events": [],
    }
    if username:
        initial_state["account_username"] = username

    _, partial_state, plan_status, resume_status, apply_status = start_partial_run(initial_state=initial_state)
    background_tasks.add_task(finish_run_background, run_id)

    return response_payload_from_state(
        partial_state,
        run_id=run_id,
        plan_status=plan_status,
        resume_status=resume_status,
        apply_status=apply_status,
    )


def _payload_from_run_entry(run_id: str, entry: dict[str, Any]) -> dict[str, Any]:
    state = entry.get("state") or {}
    return response_payload_from_state(
        state,
        run_id=run_id,
        plan_status=entry.get("status"),
        resume_status=entry.get("resume_status"),
        apply_status=entry.get("apply_status"),
        report_text=entry.get("report_text") or "",
        error=entry.get("error"),
    ) | {
        "latest_step_message": entry.get("latest_step_message") or latest_step_message(state),
    }


def _run_statuses_settled(payload: dict[str, Any]) -> bool:
    pending = (
        payload.get("plan_status") == "pending"
        or payload.get("resume_status") == "pending"
        or payload.get("apply_status") == "pending"
    )
    return not pending


def _run_has_error(payload: dict[str, Any]) -> bool:
    return payload.get("plan_status") == "error" or payload.get("apply_status") == "error"


def _sse_frame(event: str, data: dict[str, Any]) -> str:
    return f"event: {event}\ndata: {json.dumps(data, ensure_ascii=False)}\n\n"


@app.get("/api/careerpilot/result/{run_id}")
def get_careerpilot_result(run_id: str) -> dict:
    entry = run_store.get_entry(run_id)
    if not entry:
        raise HTTPException(status_code=404, detail="run_id not found")
    return _payload_from_run_entry(run_id, entry)


@app.get("/api/careerpilot/result/{run_id}/events")
def stream_careerpilot_result_events(run_id: str) -> StreamingResponse:
    """SSE progress stream: snapshot on connect, progress on each RunStore update."""
    if run_store.get_entry(run_id) is None:
        raise HTTPException(status_code=404, detail="run_id not found")

    def event_generator() -> Iterator[str]:
        q = run_store.subscribe(run_id)
        try:
            # Subscribe first, then snapshot, so we cannot miss an update between read and wait.
            current = run_store.get_entry(run_id)
            if not current:
                yield _sse_frame("error", {"ok": False, "run_id": run_id, "error": "run_id not found"})
                return

            payload = _payload_from_run_entry(run_id, current)
            if _run_has_error(payload):
                yield _sse_frame("error", payload)
                return
            if _run_statuses_settled(payload):
                yield _sse_frame("snapshot", payload)
                yield _sse_frame("done", payload)
                return

            yield _sse_frame("snapshot", payload)

            while True:
                try:
                    q.get(timeout=15.0)
                except queue.Empty:
                    yield ": ping\n\n"
                    continue

                current = run_store.get_entry(run_id)
                if not current:
                    yield _sse_frame("error", {"ok": False, "run_id": run_id, "error": "run_id not found"})
                    return

                payload = _payload_from_run_entry(run_id, current)
                if _run_has_error(payload):
                    yield _sse_frame("error", payload)
                    return
                if _run_statuses_settled(payload):
                    yield _sse_frame("progress", payload)
                    yield _sse_frame("done", payload)
                    return
                yield _sse_frame("progress", payload)
        finally:
            run_store.unsubscribe(run_id, q)

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )


@app.get("/api/careerpilot/session/{run_id}/application_pack")
def get_application_pack(
    run_id: str,
    job_id: str | None = Query(default=None),
    format: str = Query(default="json"),
) -> Any:
    fmt = (format or "json").strip().lower()
    if fmt not in ("json", "markdown", "zip"):
        raise HTTPException(status_code=400, detail="format must be json, markdown, or zip")

    entry = run_store.get_entry(run_id)
    if not entry:
        raise HTTPException(status_code=404, detail="run_id not found")

    state = dict(entry.get("state") or {})
    resume_status = str(entry.get("resume_status") or "")
    if resume_status == "pending":
        raise HTTPException(status_code=409, detail="Resume suggestions still generating")

    pack = build_application_pack(state, run_id=run_id, job_id=job_id)
    state["application_pack"] = pack
    run_store.update_state(run_id, state)

    if fmt == "json":
        return {"ok": True, "run_id": run_id, "application_pack": pack}

    if fmt == "markdown":
        body = pack_to_markdown(pack)
        fname = zip_download_filename(pack).replace(".zip", ".md")
        return Response(
            content=body.encode("utf-8"),
            media_type="text/markdown; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="{fname}"'},
        )

    zip_bytes = pack_to_zip_bytes(pack)
    fname = zip_download_filename(pack)
    return Response(
        content=zip_bytes,
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


@app.get("/api/careerpilot/users/{user_id}/memory")
def get_user_memory(user_id: str) -> dict:
    mem = load_user_memory(user_id)
    return {"ok": True, "user_id": user_id, "user_memory": mem}


@app.patch("/api/careerpilot/users/{user_id}/memory")
def patch_user_memory(user_id: str, body: UserMemoryPatch) -> dict:
    mem = load_user_memory(user_id)
    if body.preferred_target_roles is not None:
        mem["preferred_target_roles"] = body.preferred_target_roles
    if body.rejected_job_ids is not None:
        mem["rejected_job_ids"] = body.rejected_job_ids
    if body.saved_job_ids is not None:
        mem["saved_job_ids"] = body.saved_job_ids
    mem = save_user_memory(mem)
    return {"ok": True, "user_id": user_id, "user_memory": mem}


@app.post("/api/careerpilot/session/{run_id}/reject_job")
async def reject_job_for_session(run_id: str, body: RejectJobRequest) -> dict:
    entry = run_store.get_entry(run_id)
    if not entry:
        raise HTTPException(status_code=404, detail="run_id not found")

    state = dict(entry.get("state") or {})
    user_id = str(state.get("user_id") or "")
    if not user_id:
        raise HTTPException(status_code=400, detail="No user_id on this run")

    mem = load_user_memory(user_id)
    rejected = list(mem.get("rejected_job_ids") or [])
    job_id = body.job_id.strip()
    if job_id not in rejected:
        rejected.append(job_id)
    mem["rejected_job_ids"] = rejected
    mem = save_user_memory(mem)
    state["user_memory"] = mem
    run_store.update_state(run_id, state)

    return response_payload_from_state(
        state,
        run_id=run_id,
        plan_status=entry.get("status"),
        resume_status=entry.get("resume_status"),
        apply_status=entry.get("apply_status"),
        report_text=entry.get("report_text") or "",
    )


@app.post("/api/careerpilot/session/{run_id}/select_job")
async def select_job_for_session(
    run_id: str,
    body: SelectJobRequest,
    background_tasks: BackgroundTasks,
) -> dict:
    entry = run_store.get_entry(run_id)
    if not entry:
        raise HTTPException(status_code=404, detail="run_id not found")

    state = dict(entry.get("state") or {})
    matches = state.get("job_matches") or []
    job_id = body.job_id.strip()
    selected = next((j for j in matches if str(j.get("id") or "") == job_id), None)
    if not selected:
        raise HTTPException(status_code=404, detail="job_id not found in this run")

    session = dict(state.get("session_memory") or {})
    session["selected_job_id"] = job_id
    session["selected_job_title"] = str(selected.get("title") or "")
    session["updated_at"] = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
    state["session_memory"] = session

    updated_state = recompute_skill_gaps(state)
    skip_plan = bool(updated_state.get("skip_study_plan"))
    new_status = entry.get("status")

    if skip_plan:
        new_status = "skipped"
        updated_state["study_plan"] = None
    elif new_status in ("skipped", "done") and not skip_plan:
        new_status = "pending"
        updated_state["study_plan"] = None

    run_store.replace(
        run_id,
        RunSession(
            run_id=run_id,
            state=updated_state,
            status=str(new_status or "pending"),
            resume_status="pending" if settings.RESUME_OPTIMIZER_ENABLED else str(entry.get("resume_status") or "skipped"),
            apply_status="pending" if settings.APPLY_STRATEGIST_ENABLED else str(entry.get("apply_status") or "skipped"),
            report_text=entry.get("report_text") or "" if new_status == "done" else "",
            latest_step_message=latest_step_message(updated_state),
        ),
    )

    background_tasks.add_task(finish_run_background, run_id)

    return response_payload_from_state(
        updated_state,
        run_id=run_id,
        plan_status=new_status,
        resume_status="pending" if settings.RESUME_OPTIMIZER_ENABLED else entry.get("resume_status"),
        apply_status="pending" if settings.APPLY_STRATEGIST_ENABLED else entry.get("apply_status"),
        report_text=entry.get("report_text") or "" if new_status == "done" else "",
    )
